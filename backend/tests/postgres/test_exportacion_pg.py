"""E10 en PostgreSQL real: los archivos .xlsx y .pdf salen con las cifras de la empresa de la sesión (sumas
Decimal reales y la vista de ventas verdadera) y nunca mezclan a dos empresas con los mismos nombres. Las
cifras son las de `sembrar_mundo_reportes` (A con m=1: 740,00; B con m=3: 2220,00)."""

import datetime as dt
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.deps import ParametrosPeriodo
from app.core.tenant import fijar_empresa
from app.inteligencia.casos_uso.cu22_consultar_reportes_indicadores import ExportarReportesVentasInventario
from app.inteligencia.casos_uso.cu23_generar_reportes_con_ia import ExportarReporteGenerativo
from app.inteligencia.schemas import ExportarReporteRequest, FechaParcial, PeriodoSolicitado, SolicitudReporte
from app.plataforma.repository import UsuarioRepository
from app.plataforma.schemas import UsuarioCrear
from tests.comercial.ayudas_reportes_ia import PERIODO, sembrar_mundo_reportes
from tests.conftest import crear_empresa


def _administrador(db, empresa_id: int, codigo: str) -> int:
    """Un administrador de la empresa (alcance global): el caso de uso necesita un usuario real."""
    fijar_empresa(db, empresa_id)
    repo = UsuarioRepository()
    admin = repo.crear(
        db, UsuarioCrear(nombre="Admin", apellido="Ex", email=f"admin@{codigo}.example.com", password="claveSegura123")
    )
    repo.asignar_roles(db, admin, ["administrador"])
    return admin.id


def _dos_empresas(fabrica_pg):
    """(id de A, id de B, admin de A, admin de B)."""
    with fabrica_pg() as db:
        id_a = crear_empresa(db, "tienda-a").id  # los ids antes de que un commit expire las filas
        id_b = crear_empresa(db, "tienda-b").id
        sembrar_mundo_reportes(db, id_a, m=1)
        sembrar_mundo_reportes(db, id_b, m=3)
        admin_a = _administrador(db, id_a, "tienda-a")
        admin_b = _administrador(db, id_b, "tienda-b")
    return id_a, id_b, admin_a, admin_b


def _celdas(archivo) -> list:
    libro = load_workbook(BytesIO(archivo.contenido))
    return [c.value for hoja in libro.worksheets for fila in hoja.iter_rows() for c in fila if c.value is not None]


def _texto(archivo) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(BytesIO(archivo.contenido)).pages)


def _solicitud_ventas_por_sucursal() -> SolicitudReporte:
    periodo = PeriodoSolicitado(
        tipo="rango_fechas", desde=FechaParcial(dia=14, mes=9, anio=2026), hasta=FechaParcial(dia=20, mes=9, anio=2026)
    )
    return SolicitudReporte(tipo_reporte="ventas_por_sucursal", periodo=periodo)


def test_el_reporte_generativo_se_exporta_con_las_cifras_de_cada_empresa(fabrica_pg):
    id_a, id_b, admin_a, admin_b = _dos_empresas(fabrica_pg)

    for empresa_id, admin_id, total, nombre in ((id_a, admin_a, Decimal("740.00"), "tienda-a"), (id_b, admin_b, Decimal("2220.00"), "tienda-b")):
        with fabrica_pg() as db:
            fijar_empresa(db, empresa_id)
            xlsx = ExportarReporteGenerativo().ejecutar(
                db, admin_id, ExportarReporteRequest(formato="xlsx", solicitud=_solicitud_ventas_por_sucursal())
            )
            pdf = ExportarReporteGenerativo().ejecutar(
                db, admin_id, ExportarReporteRequest(formato="pdf", solicitud=_solicitud_ventas_por_sucursal())
            )

        celdas = _celdas(xlsx)
        assert nombre in celdas and float(total) in celdas, nombre
        assert xlsx.nombre_archivo.startswith(f"reporte-ventas_por_sucursal-{nombre}-") and xlsx.nombre_archivo.endswith(".xlsx")
        texto = _texto(pdf)
        assert nombre in texto and f"{total:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") in texto
        otro = "tienda-b" if nombre == "tienda-a" else "tienda-a"
        assert otro not in texto and otro not in " ".join(map(str, celdas))


def test_los_reportes_convencionales_se_exportan_con_las_cifras_de_cada_empresa(fabrica_pg):
    id_a, id_b, admin_a, admin_b = _dos_empresas(fabrica_pg)

    for empresa_id, admin_id, total in ((id_a, admin_a, 740.0), (id_b, admin_b, 2220.0)):
        with fabrica_pg() as db:
            fijar_empresa(db, empresa_id)
            cu = ExportarReportesVentasInventario()
            ventas = cu.ventas(db, admin_id, PERIODO, None, None, None, "xlsx", True)
            inventario = cu.inventario(db, admin_id, None, "xlsx", True)
            reservas = cu.reservas(db, admin_id, PERIODO, None, "pdf", True)

        assert total in _celdas(ventas)
        assert "Centro" in _celdas(inventario) and "Reporte de reservas" in _texto(reservas)


def test_un_periodo_sin_ventas_igual_exporta(fabrica_pg):
    id_a, _id_b, admin_a, _admin_b = _dos_empresas(fabrica_pg)
    enero = ParametrosPeriodo(desde=dt.date(2026, 1, 1), hasta=dt.date(2026, 1, 31))

    with fabrica_pg() as db:
        fijar_empresa(db, id_a)
        archivo = ExportarReportesVentasInventario().ventas(db, admin_a, enero, None, None, None, "pdf", True)

    assert "No hubo ventas con estos filtros." in _texto(archivo)
