"""E10 — Pruebas del constructor de Excel (`construir_xlsx`). Sin base de datos: un
DocumentoReporte es puro y el libro se lee de vuelta con openpyxl."""

import builtins
import datetime as dt
import re
import tempfile
import zipfile
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.inteligencia.exportacion_comun import (
    COLOR_ACENTO,
    COLOR_ENCABEZADO,
    EXPORTACION_MAX_FILAS_XLSX,
    PALETA_SERIES,
)
from app.inteligencia.exportacion_xlsx import construir_xlsx
from app.inteligencia.schemas import ColumnaReporte, GraficoSugerido, SeccionReporte
from tests.inteligencia.ayudas_exportacion import (
    COLUMNAS_VENTAS,
    documento,
    fila_ventas,
)


def abrir(contenido: bytes):
    return load_workbook(BytesIO(contenido))


def construir(doc, **kwargs):
    return abrir(construir_xlsx(doc, **kwargs))


def fila_encabezado(doc, *, recortada: bool = False, nota: bool = False) -> int:
    """Fila del encabezado de la tabla, calculada con la regla documentada en el módulo:
    3 líneas fijas + filtros + subtítulo + aviso de recorte + avisos + nota + 1 en blanco + 1."""
    return (
        3
        + len(doc.filtros)
        + (1 if doc.subtitulo else 0)
        + (1 if recortada else 0)
        + len(doc.avisos)
        + (1 if nota else 0)
        + 2
    )


def nombres_en_zip(contenido: bytes) -> list[str]:
    with zipfile.ZipFile(BytesIO(contenido)) as z:
        return z.namelist()


def leer_zip(contenido: bytes, nombre: str) -> str:
    with zipfile.ZipFile(BytesIO(contenido)) as z:
        return z.read(nombre).decode("utf-8")


def graficos_en(contenido: bytes) -> list[str]:
    return [n for n in nombres_en_zip(contenido) if re.fullmatch(r"xl/charts/chart\d+\.xml", n)]


# ---- Cabecera -------------------------------------------------------------------------------


def test_cabecera_titulo_empresa_y_fecha():
    doc = documento()
    ws = construir(doc).active
    assert ws["A1"].value == "Ventas por sucursal"
    assert ws["A1"].font.bold is True
    assert ws["A1"].font.size == 14
    assert ws["A2"].value == "FashionStore"
    assert ws["A3"].value == "Generado: 20/09/2026 14:05 (hora de Bolivia)"


def test_cabecera_filtros_y_subtitulo():
    doc = documento(filtros=["Periodo: 01/09/2026 al 20/09/2026", "Sucursales: todas", "Canal: tienda"])
    ws = construir(doc).active
    assert [ws.cell(row=r, column=1).value for r in (4, 5, 6)] == doc.filtros
    assert ws["A7"].value == doc.subtitulo
    assert ws.cell(row=fila_encabezado(doc) - 1, column=1).value is None  # fila en blanco


def test_cabecera_sin_filtros_ni_subtitulo_sube_el_encabezado():
    doc = documento(filtros=[], subtitulo=None)
    ws = construir(doc).active
    fila = fila_encabezado(doc)
    assert fila == 5
    assert ws.cell(row=fila, column=1).value == "Sucursal"


def test_avisos_del_documento_van_en_la_cabecera():
    doc = documento(avisos=["Aviso uno", "Aviso dos"])
    ws = construir(doc).active
    textos = [ws.cell(row=r, column=1).value for r in range(1, fila_encabezado(doc))]
    assert "Aviso uno" in textos and "Aviso dos" in textos
    assert ws.cell(row=fila_encabezado(doc), column=1).value == "Sucursal"


def test_propiedades_del_libro():
    libro = construir(documento(titulo="Ventas de septiembre"))
    assert libro.properties.title == "Ventas de septiembre"
    assert libro.properties.creator == "FashionStore"


# ---- Encabezado y datos ---------------------------------------------------------------------


def test_encabezado_de_la_tabla_con_estilo_del_diseno():
    doc = documento()
    ws = construir(doc).active
    fila = fila_encabezado(doc)
    assert [ws.cell(row=fila, column=j).value for j in range(1, 7)] == [c.titulo for c in COLUMNAS_VENTAS]
    for j in range(1, 7):
        celda = ws.cell(row=fila, column=j)
        assert celda.fill.fgColor.rgb.endswith(COLOR_ENCABEZADO)
        assert celda.font.bold is True
        assert celda.font.color.rgb.endswith("FFFFFF")
        assert celda.alignment.wrap_text is True


def test_valores_numericos_son_numeros_y_coinciden_con_las_filas():
    doc = documento(n_filas=3)
    ws = construir(doc).active
    primera = fila_encabezado(doc) + 1
    for i, fila in enumerate(doc.secciones[0].filas):
        r = primera + i
        assert ws.cell(row=r, column=1).value == fila["sucursal"]
        transacciones, total = ws.cell(row=r, column=2), ws.cell(row=r, column=3)
        assert transacciones.value == fila["transacciones"]
        assert transacciones.data_type == "n" and isinstance(transacciones.value, int)
        assert total.value == pytest.approx(fila["total_ventas"])
        assert total.data_type == "n" and isinstance(total.value, float)


def test_formatos_de_numero_por_tipo_de_columna():
    doc = documento(n_filas=1)
    ws = construir(doc).active
    r = fila_encabezado(doc) + 1
    assert ws.cell(row=r, column=2).number_format == "#,##0"
    assert ws.cell(row=r, column=3).number_format == "#,##0.00"
    assert ws.cell(row=r, column=4).number_format == "0.0%"
    assert ws.cell(row=r, column=5).number_format == "dd/mm/yyyy"


def test_porcentaje_se_guarda_como_fraccion():
    doc = documento(n_filas=1)
    ws = construir(doc).active
    assert ws.cell(row=fila_encabezado(doc) + 1, column=4).value == 0.125


def test_fecha_se_guarda_como_fecha_y_la_no_parseable_queda_texto():
    columnas = [ColumnaReporte(clave="d", titulo="Día", tipo="fecha")]
    filas = [
        {"d": "2026-09-01"},
        {"d": dt.date(2026, 9, 2)},
        {"d": "2026-09-03T10:30:00"},
        {"d": "no es fecha"},
    ]
    doc = documento(columnas=columnas, filas=filas, grafico=None)
    ws = construir(doc).active
    r = fila_encabezado(doc) + 1
    assert [ws.cell(row=r + i, column=1).value for i in range(3)] == [
        dt.datetime(2026, 9, 1),
        dt.datetime(2026, 9, 2),
        dt.datetime(2026, 9, 3),
    ]
    assert ws.cell(row=r, column=1).number_format == "dd/mm/yyyy"
    assert ws.cell(row=r + 3, column=1).value == "no es fecha"
    assert ws.cell(row=r + 3, column=1).data_type == "s"


def test_booleano_si_no_y_vacios():
    doc = documento(n_filas=2)
    doc.secciones[0].filas.append({"sucursal": "Sin datos sueltos", "transacciones": None, "total_ventas": None})
    ws = construir(doc).active
    r = fila_encabezado(doc) + 1
    assert ws.cell(row=r, column=6).value == "Sí"
    assert ws.cell(row=r + 1, column=6).value == "No"
    ultima = r + 2
    for j in range(2, 7):
        assert ws.cell(row=ultima, column=j).value is None


def test_ancho_de_columnas_entre_minimo_y_maximo():
    columnas = [
        ColumnaReporte(clave="a", titulo="X", tipo="texto"),
        ColumnaReporte(clave="b", titulo="Descripción", tipo="texto"),
    ]
    filas = [{"a": "ab", "b": "z" * 200}]
    ws = construir(documento(columnas=columnas, filas=filas, grafico=None)).active
    assert ws.column_dimensions["A"].width == 10
    assert ws.column_dimensions["B"].width == 50


def test_acentos_y_enie_intactos():
    columnas = [ColumnaReporte(clave="c", titulo="Categoría", tipo="texto")]
    doc = documento(columnas=columnas, filas=[{"c": "Ñandú"}], grafico=None, empresa="Camisería Peña")
    ws = construir(doc).active
    r = fila_encabezado(doc)
    assert ws.cell(row=r, column=1).value == "Categoría"
    assert ws.cell(row=r + 1, column=1).value == "Ñandú"
    assert ws["A2"].value == "Camisería Peña"


# ---- Totales, panel y filtro ----------------------------------------------------------------


def test_fila_de_totales_estatica_con_formato():
    doc = documento(n_filas=3, totales={"transacciones": 9, "total_ventas": 3753.0})
    ws = construir(doc).active
    r = fila_encabezado(doc) + 3 + 1
    assert ws.cell(row=r, column=1).value == "Total"
    assert ws.cell(row=r, column=1).font.bold is True
    assert ws.cell(row=r, column=2).value == 9
    assert ws.cell(row=r, column=3).value == 3753.0  # el valor dado, no una suma ni una fórmula
    assert ws.cell(row=r, column=3).number_format == "#,##0.00"
    for j in (2, 3):
        assert ws.cell(row=r, column=j).data_type == "n"  # número, nunca fórmula
    for j in (4, 5, 6):
        assert ws.cell(row=r, column=j).value is None
    assert ws.cell(row=r, column=2).border.top.style == "medium"


def test_sin_totales_no_hay_fila_de_total():
    doc = documento(n_filas=2, totales=None)
    ws = construir(doc).active
    r = fila_encabezado(doc) + 3
    assert ws.cell(row=r, column=1).value is None


def test_panel_inmovilizado_y_autofiltro_sin_la_fila_de_totales():
    doc = documento(n_filas=4, totales={"total_ventas": 1.0})
    ws = construir(doc).active
    fila = fila_encabezado(doc)
    assert ws.freeze_panes == f"A{fila + 1}"
    assert ws.auto_filter.ref == f"A{fila}:F{fila + 4}"


# ---- Gráficos -------------------------------------------------------------------------------


GRAFICOS = [
    (GraficoSugerido(tipo="barras", eje_x="sucursal", eje_y="total_ventas"), "barChart", 'barDir val="col"'),
    (GraficoSugerido(tipo="barras_horizontales", eje_x="sucursal", eje_y="total_ventas"), "barChart", 'barDir val="bar"'),
    (GraficoSugerido(tipo="linea", eje_x="sucursal", eje_y="total_ventas"), "lineChart", "<lineChart>"),
    (GraficoSugerido(tipo="torta", eje_x="sucursal", eje_y="total_ventas"), "pieChart", "<pieChart>"),
]


@pytest.mark.parametrize(("grafico", "elemento", "marca"), GRAFICOS, ids=lambda p: getattr(p, "tipo", None) or None)
def test_grafico_nativo_por_tipo(grafico, elemento, marca):
    contenido = construir_xlsx(documento(grafico=grafico))
    assert graficos_en(contenido) == ["xl/charts/chart1.xml"]
    xml = leer_zip(contenido, "xl/charts/chart1.xml")
    assert f"<{elemento}>" in xml
    assert marca in xml
    assert "<strRef><f>'Datos del grafico'!$A$2:$A$4</f></strRef></cat>" in xml  # categorías como texto
    assert "Ventas por sucursal" in xml  # título = título de la sección
    assert "Datos del grafico" in xml  # apunta a la hoja auxiliar


def test_grafico_se_ancla_a_la_derecha_de_la_tabla_a_la_altura_del_encabezado():
    doc = documento()
    contenido = construir_xlsx(doc)
    nombre = next(n for n in nombres_en_zip(contenido) if re.fullmatch(r"xl/drawings/drawing\d+\.xml", n))
    xml = leer_zip(contenido, nombre)
    columna = int(re.search(r"<col>(\d+)</col>", xml).group(1))
    fila = int(re.search(r"<row>(\d+)</row>", xml).group(1))
    assert columna == 6 + 1  # 6 columnas de tabla + una libre; 0-based => índice 7 = columna H
    assert fila + 1 == fila_encabezado(doc)


def test_colores_de_la_serie_principal_y_de_la_torta():
    barras = leer_zip(construir_xlsx(documento()), "xl/charts/chart1.xml")
    assert COLOR_ACENTO in barras
    torta = leer_zip(
        construir_xlsx(documento(grafico=GraficoSugerido(tipo="torta", eje_x="sucursal", eje_y="total_ventas"))),
        "xl/charts/chart1.xml",
    )
    for color in PALETA_SERIES[:3]:
        assert color in torta


@pytest.mark.parametrize(
    "caso",
    ["incluir_false", "tabla", "sin_filas", "sin_grafico"],
)
def test_sin_grafico_cuando_no_corresponde(caso):
    kwargs = {}
    if caso == "incluir_false":
        doc = documento()
        kwargs["incluir_grafico"] = False
    elif caso == "tabla":
        doc = documento(grafico=GraficoSugerido(tipo="tabla"))
    elif caso == "sin_filas":
        doc = documento(n_filas=0)
    else:
        doc = documento(grafico=None)
    contenido = construir_xlsx(doc, **kwargs)
    assert graficos_en(contenido) == []
    assert "Datos del grafico" not in abrir(contenido).sheetnames


def _doc_agrupado():
    columnas = [
        ColumnaReporte(clave="mes", titulo="Mes", tipo="texto"),
        ColumnaReporte(clave="canal", titulo="Canal", tipo="texto"),
        ColumnaReporte(clave="ventas", titulo="Ventas (Bs)", tipo="moneda"),
    ]
    filas = [
        {"mes": "2026-07", "canal": "Tienda", "ventas": 100.0},
        {"mes": "2026-07", "canal": "Web", "ventas": 40.0},
        {"mes": "2026-08", "canal": "Tienda", "ventas": 120.0},
        {"mes": "2026-09", "canal": "Web", "ventas": 60.0},
    ]
    grafico = GraficoSugerido(tipo="barras_agrupadas", eje_x="mes", eje_y="ventas", serie="canal")
    return documento(columnas=columnas, filas=filas, grafico=grafico)


def test_barras_agrupadas_hoja_auxiliar_con_pivote():
    contenido = construir_xlsx(_doc_agrupado())
    assert graficos_en(contenido) == ["xl/charts/chart1.xml"]
    xml = leer_zip(contenido, "xl/charts/chart1.xml")
    assert xml.count("<ser>") == 2  # una serie por canal
    libro = abrir(contenido)
    datos = libro["Datos del grafico"]
    filas = [[c.value for c in fila] for fila in datos.iter_rows(min_row=1, max_row=4, max_col=3)]
    assert filas == [
        ["Mes", "Tienda", "Web"],
        ["2026-07", 100.0, 40.0],
        ["2026-08", 120.0, 0.0],
        ["2026-09", 0.0, 60.0],
    ]
    assert COLOR_ACENTO in xml and PALETA_SERIES[1] in xml


def test_mas_de_12_categorias_se_agrupan_en_otros_en_la_hoja_auxiliar():
    filas = [fila_ventas(i, total=float(100 + i)) for i in range(20)]
    contenido = construir_xlsx(documento(filas=filas))
    datos = abrir(contenido)["Datos del grafico"]
    assert datos.max_row == 13  # encabezado + 12 categorías
    assert datos.cell(row=13, column=1).value == "Otros"
    assert datos.cell(row=13, column=2).value == pytest.approx(sum(sorted(100.0 + i for i in range(20))[:9]))
    # La tabla del reporte sigue completa.
    doc = documento(filas=filas)
    assert abrir(construir_xlsx(doc)).worksheets[0].max_row >= fila_encabezado(doc) + 20


def test_dos_secciones_con_grafico_tienen_hojas_de_datos_distintas():
    base = documento()
    segunda = SeccionReporte(
        titulo="Segunda",
        columnas=base.secciones[0].columnas,
        filas=base.secciones[0].filas,
        grafico=GraficoSugerido(tipo="linea", eje_x="sucursal", eje_y="total_ventas"),
    )
    base.secciones.append(segunda)
    contenido = construir_xlsx(base)
    assert len(graficos_en(contenido)) == 2
    libro = abrir(contenido)
    assert "Datos del grafico" in libro.sheetnames and "Datos del grafico 2" in libro.sheetnames
    assert "Datos del grafico 2" in leer_zip(contenido, "xl/charts/chart2.xml")


# ---- Inyección de fórmulas ------------------------------------------------------------------

PELIGROSOS = ['=HYPERLINK("http://x.test","clic")', "+cmd|' /C calc'!A0", "-1+1", "@SUM(A1:A2)", "  =1+1", "\t=1+1"]


@pytest.mark.parametrize("peligroso", PELIGROSOS)
def test_valores_que_parecen_formula_quedan_como_texto(peligroso):
    columnas = [ColumnaReporte(clave="c", titulo=peligroso, tipo="texto")]
    doc = documento(
        columnas=columnas,
        filas=[{"c": peligroso}],
        grafico=None,
        titulo=peligroso,
        empresa=peligroso,
        filtros=[peligroso],
        subtitulo=peligroso,
        avisos=[peligroso],
    )
    contenido = construir_xlsx(doc)
    ws = abrir(contenido).active
    celdas = [ws["A1"], ws["A2"], ws["A4"], ws["A5"], ws["A6"]]
    fila = fila_encabezado(doc)
    celdas += [ws.cell(row=fila, column=1), ws.cell(row=fila + 1, column=1)]
    for celda in celdas:
        assert celda.data_type == "s", celda.coordinate
        assert celda.value.strip() == peligroso.strip()
    assert "<f>" not in leer_zip(contenido, "xl/worksheets/sheet1.xml")


def test_texto_en_columna_numerica_y_codigos_de_error_siguen_siendo_texto():
    columnas = [
        ColumnaReporte(clave="m", titulo="Monto", tipo="moneda"),
        ColumnaReporte(clave="t", titulo="Nota", tipo="texto"),
    ]
    filas = [{"m": "=1+1", "t": "#REF!"}, {"m": "n/d", "t": "#N/A"}]
    doc = documento(columnas=columnas, filas=filas, grafico=None)
    ws = abrir(construir_xlsx(doc)).active
    r = fila_encabezado(doc) + 1
    for i in range(2):
        for j in (1, 2):
            assert ws.cell(row=r + i, column=j).data_type == "s"


def test_totales_no_son_formulas():
    doc = documento(n_filas=2, totales={"total_ventas": 2501.0})
    contenido = construir_xlsx(doc)
    assert "<f>" not in leer_zip(contenido, "xl/worksheets/sheet1.xml")


# ---- Nombres de hoja ------------------------------------------------------------------------


def test_nombre_de_hoja_largo_o_con_caracteres_invalidos_se_sanea():
    titulo = "Ventas: [2026] / Sucursal*Centro? \\ con un nombre larguísimo"
    doc = documento(titulo=titulo)
    nombre = construir(doc).sheetnames[0]
    assert 0 < len(nombre) <= 31
    assert not set(nombre) & set(":\\/?*[]")
    assert nombre.startswith("Ventas")


def test_nombre_de_hoja_vacio_usa_reporte():
    doc = documento(titulo="???")
    assert construir(doc).sheetnames[0] == "Reporte"


def test_dos_secciones_con_el_mismo_titulo_tienen_hojas_distintas():
    doc = documento(grafico=None)
    doc.secciones.append(doc.secciones[0].model_copy())
    nombres = construir(doc).sheetnames
    assert len(nombres) == 2
    assert len({n.lower() for n in nombres}) == 2
    assert all(len(n) <= 31 for n in nombres)


def test_seccion_con_nombre_largo_repetido_sigue_dentro_de_31_caracteres():
    doc = documento(grafico=None, titulo="x" * 60)
    doc.secciones.append(doc.secciones[0].model_copy())
    doc.secciones.append(doc.secciones[0].model_copy())
    nombres = construir(doc).sheetnames
    assert len({n.lower() for n in nombres}) == 3 and all(len(n) <= 31 for n in nombres)


# ---- Recortes y tabla vacía -----------------------------------------------------------------


def test_max_filas_recorta_y_avisa():
    doc = documento(n_filas=8, grafico=None)
    ws = construir(doc, max_filas=5).active
    fila = fila_encabezado(doc, recortada=True)
    textos = [ws.cell(row=r, column=1).value for r in range(4, fila)]
    assert "Se muestran las primeras 5 filas" in textos
    assert ws.cell(row=fila + 5, column=1).value == "Sucursal 004"
    assert ws.cell(row=fila + 6, column=1).value is None
    assert ws.auto_filter.ref == f"A{fila}:F{fila + 5}"


def test_max_filas_tambien_recorta_los_datos_del_grafico():
    doc = documento(n_filas=8)
    datos = construir(doc, max_filas=5)["Datos del grafico"]
    assert datos.max_row == 6


def test_documento_truncado_muestra_aviso_con_las_filas_escritas():
    doc = documento(n_filas=3, truncado=True, grafico=None)
    ws = construir(doc).active
    fila = fila_encabezado(doc, recortada=True)
    textos = [ws.cell(row=r, column=1).value for r in range(4, fila)]
    assert "Se muestran las primeras 3 filas" in textos


def test_tabla_vacia_tiene_encabezado_y_sin_datos():
    doc = documento(n_filas=0, totales={"total_ventas": 0})
    contenido = construir_xlsx(doc)
    ws = abrir(contenido).active
    fila = fila_encabezado(doc)
    assert ws.cell(row=fila, column=1).value == "Sucursal"
    assert ws.cell(row=fila + 1, column=1).value == "Sin datos"
    assert ws.cell(row=fila + 2, column=1).value is None  # sin totales
    assert graficos_en(contenido) == []


def test_documento_sin_secciones_es_un_libro_valido():
    doc = documento()
    doc.secciones.clear()
    libro = construir(doc)
    assert libro.sheetnames == ["Reporte"]
    assert libro.active["A1"].value == "Ventas por sucursal"


def test_5001_filas_se_recortan_a_5000_y_avisan():
    filas = [fila_ventas(i) for i in range(EXPORTACION_MAX_FILAS_XLSX + 1)]
    doc = documento(filas=filas, grafico=None)
    ws = construir(doc).active
    fila = fila_encabezado(doc, recortada=True)
    assert ws.cell(row=fila + 5000, column=1).value == "Sucursal 4999"
    assert ws.cell(row=fila + 5001, column=1).value is None
    textos = [ws.cell(row=r, column=1).value for r in range(4, fila)]
    assert "Se muestran las primeras 5000 filas" in textos


# ---- Memoria, no disco ----------------------------------------------------------------------


def test_no_escribe_en_disco(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))

    abrir_real = builtins.open

    def abrir_solo_lectura(archivo, modo="r", *args, **kwargs):
        if any(c in str(modo) for c in "wax+"):
            raise AssertionError(f"Se intentó abrir {archivo!r} en modo {modo!r}")
        return abrir_real(archivo, modo, *args, **kwargs)

    def sin_temporales(*args, **kwargs):
        raise AssertionError("Se intentó crear un archivo temporal")

    monkeypatch.setattr(builtins, "open", abrir_solo_lectura)
    monkeypatch.setattr(tempfile, "mkstemp", sin_temporales)
    monkeypatch.setattr(tempfile, "NamedTemporaryFile", sin_temporales)
    monkeypatch.setattr(tempfile, "TemporaryFile", sin_temporales)

    contenido = construir_xlsx(_doc_agrupado())
    assert contenido[:2] == b"PK"
    assert list(tmp_path.iterdir()) == []


def test_devuelve_bytes_de_un_zip_xlsx_valido():
    contenido = construir_xlsx(documento())
    assert isinstance(contenido, bytes)
    assert "xl/workbook.xml" in nombres_en_zip(contenido)
