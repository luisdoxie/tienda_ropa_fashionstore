"""E10 — Constructor de Excel (.xlsx) con openpyxl.

`construir_xlsx` recibe un DocumentoReporte y devuelve los bytes del libro. Todo ocurre en memoria
(BytesIO): no se escribe ningún archivo ni se crean temporales.

Diseño de cada hoja (una por sección):

    fila 1            título del documento (negrita 14)
    fila 2            nombre de la empresa
    fila 3            "Generado: dd/mm/aaaa hh:mm (hora de Bolivia)"
    filas siguientes  una por cada filtro de `doc.filtros`
                      el subtítulo (si hay)
                      el aviso "Se muestran las primeras N filas" (si se recortó o `doc.truncado`)
                      un renglón por cada aviso de `doc.avisos`
                      la nota de la sección (si hay)
                      una fila en blanco
    fila E            ENCABEZADO DE LA TABLA

La fila E del encabezado NO es fija: depende de cuántos filtros, avisos y demás líneas hay. Vale
`3 + len(filtros) + [subtítulo] + [aviso de recorte] + len(avisos) + [nota] + 1 (en blanco) + 1`,
donde cada [..] suma 1 si existe. Los datos empiezan en E + 1; el panel inmovilizado queda en
`A{E+1}`; la fila de totales (si hay) va justo después del último dato; el gráfico se ancla en la
fila E, una columna libre a la derecha de la tabla.

Seguridad: todo texto que venga de los datos se escribe como cadena explícita (`data_type = "s"`),
de modo que un valor como `=HYPERLINK(...)`, `+cmd`, `-1+1` o `@SUM(A1)` nunca es una fórmula.
Los totales son valores estáticos, nunca fórmulas.

El gráfico es nativo de Excel. Sus datos se escriben en una hoja auxiliar "Datos del grafico"
(una por gráfico: "Datos del grafico", "Datos del grafico 2", ...) con las categorías en la
primera columna y una columna por serie, ya recortadas a 12 categorías ("Otros") por
`exportacion_comun.datos_de_grafico`.
"""

import datetime as dt
import re
from decimal import Decimal
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.writer.excel import ExcelWriter

from app.inteligencia.exportacion_comun import (
    COLOR_ACENTO,
    COLOR_BORDE,
    COLOR_ENCABEZADO,
    COLOR_FILA_ALTERNA,
    COLOR_TEXTO_ENCABEZADO,
    EXPORTACION_MAX_FILAS_XLSX,
    PALETA_SERIES,
    DatosGrafico,
    a_decimal,
    datos_de_grafico,
    texto_de_generacion,
)
from app.inteligencia.schemas import DocumentoReporte, SeccionReporte

FORMATOS = {
    "moneda": "#,##0.00",
    "entero": "#,##0",
    "porcentaje": "0.0%",
    "fecha": "dd/mm/yyyy",
}
NOMBRE_HOJA_RESPALDO = "Reporte"
NOMBRE_HOJA_DATOS = "Datos del grafico"
MAX_LARGO_HOJA = 31
MAX_LARGO_CELDA = 32767
ANCHO_MIN, ANCHO_MAX = 10, 50
CREADOR = "FashionStore"
TEXTO_BOOL = {True: "Sí", False: "No"}

_HOJA_INVALIDOS = re.compile(r"[:\\/?*\[\]]")
# Caracteres de control que XML no admite (openpyxl los rechaza); se quitan, salvo tab, LF y CR.
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

_LADO = Side(style="thin", color=COLOR_BORDE)
_BORDE = Border(left=_LADO, right=_LADO, top=_LADO, bottom=_LADO)
_BORDE_TOTAL = Border(left=_LADO, right=_LADO, bottom=_LADO, top=Side(style="medium", color=COLOR_ENCABEZADO))
_RELLENO_ENCABEZADO = PatternFill("solid", start_color=COLOR_ENCABEZADO, end_color=COLOR_ENCABEZADO)
_RELLENO_ALTERNO = PatternFill("solid", start_color=COLOR_FILA_ALTERNA, end_color=COLOR_FILA_ALTERNA)
_FUENTE_ENCABEZADO = Font(bold=True, color=COLOR_TEXTO_ENCABEZADO)
_FUENTE_TITULO = Font(bold=True, size=14)
_FUENTE_NEGRITA = Font(bold=True)


# ---- Texto y valores ------------------------------------------------------------------------


def _limpiar(valor) -> str:
    texto = _CONTROL.sub("", str(valor))
    return texto[:MAX_LARGO_CELDA]


def _poner_texto(celda, valor) -> int:
    """Escribe `valor` como cadena explícita (nunca fórmula ni error). Devuelve su largo."""
    texto = _limpiar(valor)
    celda.value = texto
    celda.data_type = "s"
    return len(texto)


def _fecha_de(valor) -> dt.date | None:
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    if isinstance(valor, str):
        try:
            return dt.date.fromisoformat(valor.strip()[:10])
        except ValueError:
            return None
    return None


def _poner_valor(celda, valor, tipo: str) -> int:
    """Escribe `valor` según el tipo de columna y devuelve el largo aproximado de lo mostrado."""
    if valor is None:
        return 0
    if isinstance(valor, bool):
        return _poner_texto(celda, TEXTO_BOOL[valor])
    formato = FORMATOS.get(tipo)
    if tipo == "fecha":
        fecha = _fecha_de(valor)
        if fecha is None:
            return _poner_texto(celda, valor)
        celda.value = fecha
        celda.number_format = formato
        return 10
    if tipo in ("moneda", "entero", "porcentaje"):
        numero = a_decimal(valor)
        if numero is None or not numero.is_finite():
            return _poner_texto(celda, valor)
        if tipo == "porcentaje":
            # Las filas traen 0..100 (12.5 = 12,5 %); Excel espera la fracción 0.125.
            numero = numero / Decimal(100)
            celda.value = float(numero)
        elif tipo == "entero" and numero == numero.to_integral_value():
            celda.value = int(numero)
        else:
            celda.value = float(numero)
        celda.number_format = formato
        return len(f"{float(celda.value):,.2f}") + (1 if tipo == "porcentaje" else 0)
    return _poner_texto(celda, valor)


# ---- Hojas ----------------------------------------------------------------------------------


def _nombre_hoja(base: str, usados: set[str]) -> str:
    """Nombre válido de hoja: sin :\\/?*[], máximo 31 caracteres, no vacío y único (sin importar
    mayúsculas) dentro del libro."""
    limpio = _CONTROL.sub("", _HOJA_INVALIDOS.sub(" ", base or "")).strip().strip("'").strip()
    limpio = limpio[:MAX_LARGO_HOJA].strip().strip("'").strip() or NOMBRE_HOJA_RESPALDO
    nombre, n = limpio, 1
    while nombre.lower() in usados:
        n += 1
        sufijo = f" {n}"
        nombre = limpio[: MAX_LARGO_HOJA - len(sufijo)].rstrip() + sufijo
    usados.add(nombre.lower())
    return nombre


def _lineas_de_cabecera(doc: DocumentoReporte, seccion: SeccionReporte | None, escritas: int, recortada: bool) -> list[str]:
    lineas = list(doc.filtros)
    if doc.subtitulo:
        lineas.append(doc.subtitulo)
    if recortada:
        lineas.append(f"Se muestran las primeras {escritas} filas")
    lineas.extend(doc.avisos)
    if seccion is not None and seccion.nota:
        lineas.append(seccion.nota)
    return lineas


def _escribir_seccion(ws, doc: DocumentoReporte, seccion: SeccionReporte | None, max_filas: int) -> dict:
    """Escribe una hoja. Devuelve {'fila_encabezado', 'n_columnas'} para anclar el gráfico."""
    columnas = seccion.columnas if seccion is not None else []
    todas = seccion.filas if seccion is not None else []
    filas = todas[: max(max_filas, 0)]
    recortada = doc.truncado or len(todas) > len(filas)

    celda = ws.cell(row=1, column=1)
    _poner_texto(celda, doc.titulo)
    celda.font = _FUENTE_TITULO
    _poner_texto(ws.cell(row=2, column=1), doc.empresa_nombre)
    _poner_texto(ws.cell(row=3, column=1), "Generado: " + texto_de_generacion(doc.generado_en))
    fila = 4
    for linea in _lineas_de_cabecera(doc, seccion, len(filas), recortada):
        _poner_texto(ws.cell(row=fila, column=1), linea)
        fila += 1
    fila_encabezado = fila + 1  # una fila en blanco antes de la tabla

    anchos = [ANCHO_MIN] * len(columnas)
    for j, columna in enumerate(columnas, start=1):
        celda = ws.cell(row=fila_encabezado, column=j)
        _poner_texto(celda, columna.titulo)
        celda.fill = _RELLENO_ENCABEZADO
        celda.font = _FUENTE_ENCABEZADO
        celda.border = _BORDE
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        anchos[j - 1] = max(anchos[j - 1], min(len(columna.titulo), ANCHO_MAX))

    ultima = fila_encabezado
    if columnas and not filas:
        ultima += 1
        celda = ws.cell(row=ultima, column=1)
        _poner_texto(celda, "Sin datos")
        celda.border = _BORDE
        for j in range(2, len(columnas) + 1):
            ws.cell(row=ultima, column=j).border = _BORDE
    else:
        for i, datos in enumerate(filas, start=1):
            ultima = fila_encabezado + i
            for j, columna in enumerate(columnas, start=1):
                celda = ws.cell(row=ultima, column=j)
                largo = _poner_valor(celda, datos.get(columna.clave), columna.tipo)
                celda.border = _BORDE
                if i % 2 == 0:
                    celda.fill = _RELLENO_ALTERNO
                if largo + 2 > anchos[j - 1]:
                    anchos[j - 1] = min(largo + 2, ANCHO_MAX)

    if filas and seccion is not None and seccion.totales is not None and columnas:
        ultima_dato = ultima
        ultima += 1
        for j, columna in enumerate(columnas, start=1):
            celda = ws.cell(row=ultima, column=j)
            celda.font = _FUENTE_NEGRITA
            celda.border = _BORDE_TOTAL
            if j == 1:
                _poner_texto(celda, "Total")
            elif columna.clave in seccion.totales:
                largo = _poner_valor(celda, seccion.totales[columna.clave], columna.tipo)
                if largo + 2 > anchos[j - 1]:
                    anchos[j - 1] = min(largo + 2, ANCHO_MAX)
        ultima_filtro = ultima_dato
    else:
        ultima_filtro = ultima

    for j, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(ANCHO_MIN, min(ANCHO_MAX, ancho))
    if columnas:
        ws.freeze_panes = f"A{fila_encabezado + 1}"
        ws.auto_filter.ref = f"A{fila_encabezado}:{get_column_letter(len(columnas))}{ultima_filtro}"
    return {"fila_encabezado": fila_encabezado, "n_columnas": len(columnas)}


# ---- Gráfico --------------------------------------------------------------------------------


def _escribir_datos_grafico(ws_datos, datos: DatosGrafico) -> None:
    """Categorías en la columna A y una columna por serie, con la fila 1 de encabezados."""
    _poner_texto(ws_datos.cell(row=1, column=1), datos.titulo_categorias or "Categoría")
    for j, nombre in enumerate(datos.series, start=2):
        _poner_texto(ws_datos.cell(row=1, column=j), nombre or "Serie")
    for i, categoria in enumerate(datos.categorias, start=2):
        _poner_texto(ws_datos.cell(row=i, column=1), categoria)
    for j, valores in enumerate(datos.series.values(), start=2):
        for i, valor in enumerate(valores, start=2):
            ws_datos.cell(row=i, column=j).value = valor
    ws_datos.column_dimensions["A"].width = 24
    for j in range(2, len(datos.series) + 2):
        ws_datos.column_dimensions[get_column_letter(j)].width = 18


def _color_de_serie(indice: int) -> str:
    return COLOR_ACENTO if indice == 0 else PALETA_SERIES[indice % len(PALETA_SERIES)]


def _crear_grafico(ws_datos, datos: DatosGrafico, titulo: str):
    n_series = len(datos.series)
    n_filas = len(datos.categorias) + 1
    if datos.tipo == "torta":
        grafico = PieChart()
    elif datos.tipo == "linea":
        grafico = LineChart()
    else:
        grafico = BarChart()
        grafico.type = "bar" if datos.tipo == "barras_horizontales" else "col"
        grafico.grouping = "clustered"
        grafico.gapWidth = 60

    grafico.title = _limpiar(titulo) or None
    valores = Reference(ws_datos, min_col=2, max_col=n_series + 1, min_row=1, max_row=n_filas)
    categorias = Reference(ws_datos, min_col=1, min_row=2, max_row=n_filas)
    grafico.add_data(valores, titles_from_data=True)
    grafico.set_categories(categorias)
    for serie in grafico.series:
        # Las categorías son texto: referencia de cadenas (openpyxl las deja como numéricas).
        serie.cat = AxDataSource(strRef=StrRef(f=str(categorias)))

    for i, serie in enumerate(grafico.series):
        color = _color_de_serie(i)
        if datos.tipo == "linea":
            serie.graphicalProperties.line.solidFill = color
            serie.graphicalProperties.line.width = 28575
            serie.marker.symbol = "circle"
            serie.marker.size = 6
            serie.marker.graphicalProperties.solidFill = color
            serie.marker.graphicalProperties.line.solidFill = color
            serie.smooth = False
        elif datos.tipo == "torta":
            for k in range(len(datos.categorias)):
                punto = DataPoint(idx=k)
                punto.graphicalProperties.solidFill = PALETA_SERIES[k % len(PALETA_SERIES)]
                serie.dPt.append(punto)
        else:
            serie.graphicalProperties.solidFill = color
            serie.graphicalProperties.line.solidFill = color

    if datos.tipo != "torta":
        grafico.x_axis.title = _limpiar(datos.titulo_categorias) or None
        grafico.y_axis.title = _limpiar(datos.titulo_valores) or None
        # openpyxl 3.1 deja los ejes ocultos en Excel reciente si no se piden explícitamente.
        grafico.x_axis.delete = False
        grafico.y_axis.delete = False
    if datos.tipo == "torta":
        grafico.legend.position = "r"
    elif n_series > 1:
        grafico.legend.position = "b"
    else:
        grafico.legend = None
    grafico.width = 18
    grafico.height = 9
    return grafico


# ---- Punto de entrada -----------------------------------------------------------------------


class _EscritorEnMemoria(ExcelWriter):
    """openpyxl vuelca cada hoja a un archivo temporal en disco antes de comprimirla. Este
    escritor serializa la hoja en un BytesIO: así `construir_xlsx` no toca el disco."""

    def write_worksheet(self, ws):
        ws._drawing = SpreadsheetDrawing()
        ws._drawing.charts = ws._charts
        ws._drawing.images = ws._images
        escritor = WorksheetWriter(ws, out=BytesIO())
        escritor.write()
        ws._rels = escritor._rels
        self._archive.writestr(ws.path[1:], escritor.read())
        self.manifest.append(ws)


def _guardar_en_memoria(libro: Workbook) -> bytes:
    salida = BytesIO()
    libro.properties.modified = dt.datetime.now(tz=dt.timezone.utc).replace(tzinfo=None)
    _EscritorEnMemoria(libro, ZipFile(salida, "w", ZIP_DEFLATED, allowZip64=True)).save()
    return salida.getvalue()


def construir_xlsx(
    doc: DocumentoReporte, incluir_grafico: bool = True, max_filas: int = EXPORTACION_MAX_FILAS_XLSX
) -> bytes:
    """El libro de Excel del documento, en bytes. Una hoja por sección (ver el docstring del
    módulo para la disposición y la fila del encabezado de la tabla). Si una sección trae más de
    `max_filas` filas se escriben solo las primeras y la hoja lo avisa."""
    libro = Workbook()
    usados: set[str] = set()
    secciones: list[SeccionReporte | None] = list(doc.secciones) or [None]
    hojas = []
    for indice, seccion in enumerate(secciones):
        ws = libro.active if indice == 0 else libro.create_sheet()
        ws.title = _nombre_hoja(seccion.titulo if seccion is not None else NOMBRE_HOJA_RESPALDO, usados)
        info = _escribir_seccion(ws, doc, seccion, max_filas)
        hojas.append((ws, seccion, info))

    if incluir_grafico:
        n_datos = 0
        for ws, seccion, info in hojas:
            datos = datos_de_grafico(seccion) if seccion is not None else None
            if datos is None or info["n_columnas"] == 0:
                continue
            # Los datos del gráfico salen de las filas ya recortadas a `max_filas`.
            if len(seccion.filas) > max_filas:
                datos = datos_de_grafico(seccion.model_copy(update={"filas": seccion.filas[: max(max_filas, 0)]}))
                if datos is None:
                    continue
            n_datos += 1
            base = NOMBRE_HOJA_DATOS if n_datos == 1 else f"{NOMBRE_HOJA_DATOS} {n_datos}"
            ws_datos = libro.create_sheet(_nombre_hoja(base, usados))
            _escribir_datos_grafico(ws_datos, datos)
            grafico = _crear_grafico(ws_datos, datos, seccion.titulo)
            ancla = f"{get_column_letter(info['n_columnas'] + 2)}{info['fila_encabezado']}"
            ws.add_chart(grafico, ancla)

    libro.properties.title = _limpiar(doc.titulo)
    libro.properties.creator = CREADOR
    return _guardar_en_memoria(libro)
